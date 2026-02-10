"""
Test script for Data Bank extraction service and API endpoints.
Creates sample Excel files and tests the full upload pipeline.
"""
import os
import sys

# Ensure we can import the app
sys.path.insert(0, os.path.dirname(__file__))
os.chdir(os.path.dirname(__file__))

import openpyxl
from pathlib import Path


def create_sample_sales_comps(filepath: str):
    """Create a sample sales comp Excel file for testing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Comps"

    # Headers
    headers = [
        "Property Name", "Market", "Submarket", "State", "Address",
        "Property Type", "Sale Date", "Year Built", "Units",
        "Sale Price", "Price Per Unit", "Cap Rate", "Occupancy",
        "Buyer", "Seller",
    ]
    ws.append(headers)

    # Sample data
    data = [
        ["Oakwood Apartments", "Dallas-Fort Worth", "Uptown", "TX", "123 Main St, Dallas, TX 75201",
         "Multifamily", "03/15/2025", "1998/2019", 240,
         "$63.88M", 266167, "5.25%", "94%",
         "Greystar", "Blackstone"],
        ["Pine Valley Gardens", "Dallas-Fort Worth", "Richardson", "TX", "456 Oak Ave, Richardson, TX 75080",
         "Multifamily", "01/22/2025", 2005, 180,
         48600000, 270000, 0.0485, 0.97,
         "Starwood Capital", "Morgan Properties"],
        ["Riverside Commons", "Dallas-Fort Worth", "Las Colinas", "TX", "789 River Rd, Irving, TX 75039",
         "Multifamily", "06/01/2024", 2015, 320,
         "$92.5M", "$289,062", "4.75%", "96%",
         "CBRE Investment", "Hines"],
        ["Legacy at Plano", "Dallas-Fort Worth", "Plano", "TX", "321 Legacy Dr, Plano, TX 75024",
         "Multifamily", "11/30/2024", "2001", 150,
         39000000, 260000, "5.1%", "93%",
         "TBD", "Lincoln Property"],
        ["Metro at Addison", "Dallas-Fort Worth", "Addison", "TX", "555 Belt Line Rd, Addison, TX 75001",
         "Multifamily", "09/10/2024", 2010, 200,
         58000000, 290000, "4.9%", "95%",
         "Crow Holdings", "Invesco"],
    ]

    for row in data:
        ws.append(row)

    wb.save(filepath)
    print(f"Created sample sales comps: {filepath}")


def create_sample_pipeline(filepath: str):
    """Create a sample pipeline tracker Excel file for testing."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pipeline Tracker"

    headers = [
        "Project Name", "Address", "County", "Metro", "Submarket",
        "Units", "Status", "Developer", "Expected Delivery", "Property Type",
    ]
    ws.append(headers)

    data = [
        ["The Vue at Victory Park", "2500 Victory Ave", "Dallas", "Dallas-Fort Worth", "Uptown",
         350, "Under Construction", "Trammell Crow", "Q3 2026", "Multifamily"],
        ["Addison Quarter", "4800 Belt Line Rd", "Dallas", "Dallas-Fort Worth", "Addison",
         220, "Lease-Up", "JPI", "Q1 2025", "Multifamily"],
        ["Legacy West Phase III", "7800 Windrose Ave", "Collin", "Dallas-Fort Worth", "Plano",
         400, "Proposed", "KDC", "Q4 2027", "Multifamily"],
        ["Las Colinas Urban Center", "201 E Las Colinas Blvd", "Dallas", "Dallas-Fort Worth", "Las Colinas",
         275, "Under Construction", "Hines", "Q2 2026", "Multifamily"],
        ["Richardson Innovation Quarter", "1000 E Campbell Rd", "Dallas", "Dallas-Fort Worth", "Richardson",
         180, "Planning", "Lincoln Property", "Q1 2028", "Multifamily"],
    ]

    for row in data:
        ws.append(row)

    wb.save(filepath)
    print(f"Created sample pipeline tracker: {filepath}")


def test_document_type_detection():
    """Test the document type detection logic."""
    from app.services.data_bank_extraction_service import detect_document_type

    comps_path = "/tmp/test_sales_comps.xlsx"
    pipeline_path = "/tmp/test_pipeline.xlsx"

    create_sample_sales_comps(comps_path)
    create_sample_pipeline(pipeline_path)

    comp_type = detect_document_type(comps_path)
    print(f"\nSales comps detected as: {comp_type}")
    assert comp_type == "sales_comps", f"Expected 'sales_comps', got '{comp_type}'"

    pipeline_type = detect_document_type(pipeline_path)
    print(f"Pipeline tracker detected as: {pipeline_type}")
    assert pipeline_type == "pipeline_tracker", f"Expected 'pipeline_tracker', got '{pipeline_type}'"

    print("Document type detection: PASS")


def test_extraction_service():
    """Test the full extraction pipeline (Python fallback mode, no Claude API)."""
    from app.database import engine, Base, SessionLocal
    from app.models.data_bank import DataBankDocument, SalesComp, PipelineProject
    from app.models.user import User
    from app.services.data_bank_extraction_service import process_data_bank_upload

    # Create tables
    Base.metadata.create_all(bind=engine)

    db = SessionLocal()
    try:
        # Create a test user if not exists
        test_user = db.query(User).filter(User.email == "test@astra.dev").first()
        if not test_user:
            from passlib.context import CryptContext
            pwd_ctx = CryptContext(schemes=["bcrypt"], deprecated="auto")
            test_user = User(
                email="test@astra.dev",
                hashed_password=pwd_ctx.hash("testpass123"),
                full_name="Test User",
            )
            db.add(test_user)
            db.commit()
            db.refresh(test_user)
            print(f"\nCreated test user: {test_user.id}")
        else:
            print(f"\nUsing existing test user: {test_user.id}")

        user_id = str(test_user.id)

        # Test sales comp extraction
        comps_path = "/tmp/test_sales_comps.xlsx"
        create_sample_sales_comps(comps_path)

        doc = DataBankDocument(
            user_id=user_id,
            filename="test_sales_comps.xlsx",
            file_path=comps_path,
            document_type="unknown",
            extraction_status="processing",
        )
        db.add(doc)
        db.flush()

        doc_type, count, warnings = process_data_bank_upload(comps_path, user_id, db, doc)
        print(f"\nSales comp extraction:")
        print(f"  Type: {doc_type}")
        print(f"  Records: {count}")
        print(f"  Warnings: {warnings}")
        print(f"  Status: {doc.extraction_status}")
        assert doc_type == "sales_comps"
        assert count == 5
        assert doc.extraction_status == "completed"

        # Verify records in DB
        comps = db.query(SalesComp).filter(SalesComp.document_id == doc.id).all()
        print(f"  DB records: {len(comps)}")
        for c in comps:
            print(f"    - {c.property_name}: cap_rate={c.cap_rate}, sale_price={c.sale_price}, units={c.units}")

        # Test pipeline extraction
        pipeline_path = "/tmp/test_pipeline.xlsx"
        create_sample_pipeline(pipeline_path)

        doc2 = DataBankDocument(
            user_id=user_id,
            filename="test_pipeline.xlsx",
            file_path=pipeline_path,
            document_type="unknown",
            extraction_status="processing",
        )
        db.add(doc2)
        db.flush()

        doc_type2, count2, warnings2 = process_data_bank_upload(pipeline_path, user_id, db, doc2)
        print(f"\nPipeline extraction:")
        print(f"  Type: {doc_type2}")
        print(f"  Records: {count2}")
        print(f"  Warnings: {warnings2}")
        print(f"  Status: {doc2.extraction_status}")
        assert doc_type2 == "pipeline_tracker"
        assert count2 == 5
        assert doc2.extraction_status == "completed"

        projects = db.query(PipelineProject).filter(PipelineProject.document_id == doc2.id).all()
        print(f"  DB records: {len(projects)}")
        for p in projects:
            print(f"    - {p.project_name}: status={p.status}, units={p.units}")

        print("\nExtraction service tests: PASS")

    finally:
        db.close()


def test_api_endpoints():
    """Test the API endpoints using the FastAPI test client."""
    from app.database import engine, Base
    from app.main import app
    from fastapi.testclient import TestClient

    # Create all tables
    Base.metadata.create_all(bind=engine)

    client = TestClient(app)

    # Login to get auth token
    login_resp = client.post("/api/v1/auth/login", json={
        "email": "test@astra.dev",
        "password": "testpass123",
    })
    print(f"\nLogin response: {login_resp.status_code}")

    if login_resp.status_code != 200:
        print(f"  Login failed: {login_resp.text}")
        print("  Creating user first...")
        reg_resp = client.post("/api/v1/auth/register", json={
            "email": "test@astra.dev",
            "password": "testpass123",
            "full_name": "Test User",
        })
        print(f"  Register response: {reg_resp.status_code}")
        login_resp = client.post("/api/v1/auth/login", json={
            "email": "test@astra.dev",
            "password": "testpass123",
        })
        print(f"  Login retry: {login_resp.status_code}")

    token = login_resp.json().get("access_token")
    if not token:
        print(f"  Could not get token. Response: {login_resp.json()}")
        return

    headers = {"Authorization": f"Bearer {token}"}

    # Test upload - sales comps
    comps_path = "/tmp/test_sales_comps.xlsx"
    create_sample_sales_comps(comps_path)

    with open(comps_path, "rb") as f:
        upload_resp = client.post(
            "/api/v1/data-bank/upload",
            files={"file": ("test_sales_comps.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )
    print(f"\nUpload sales comps: {upload_resp.status_code}")
    print(f"  Response: {upload_resp.json()}")
    assert upload_resp.status_code == 201
    doc_id = upload_resp.json()["document_id"]

    # Test upload - pipeline
    pipeline_path = "/tmp/test_pipeline.xlsx"
    create_sample_pipeline(pipeline_path)

    with open(pipeline_path, "rb") as f:
        upload_resp2 = client.post(
            "/api/v1/data-bank/upload",
            files={"file": ("test_pipeline.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )
    print(f"\nUpload pipeline: {upload_resp2.status_code}")
    print(f"  Response: {upload_resp2.json()}")
    assert upload_resp2.status_code == 201

    # Test list documents
    list_resp = client.get("/api/v1/data-bank/documents", headers=headers)
    print(f"\nList documents: {list_resp.status_code}")
    print(f"  Total: {list_resp.json()['total']}")
    for doc in list_resp.json()["documents"]:
        print(f"  - [{doc['document_type']}] {doc['filename']} ({doc['record_count']} records, {doc['extraction_status']})")
    assert list_resp.status_code == 200

    # Test get document detail
    detail_resp = client.get(f"/api/v1/data-bank/document/{doc_id}", headers=headers)
    print(f"\nGet document {doc_id}: {detail_resp.status_code}")
    print(f"  Type: {detail_resp.json()['document_type']}")
    assert detail_resp.status_code == 200

    # Test query comps (should return the extracted comps)
    comps_resp = client.get("/api/v1/data-bank/comps", headers=headers)
    print(f"\nQuery comps: {comps_resp.status_code}")
    print(f"  Comps found: {len(comps_resp.json())}")
    for c in comps_resp.json():
        print(f"  - {c['property_name']}: cap={c['cap_rate']}, price={c['sale_price']}, units={c['units']}")
    assert comps_resp.status_code == 200

    # Test query pipeline
    pipeline_resp = client.get("/api/v1/data-bank/pipeline", headers=headers)
    print(f"\nQuery pipeline: {pipeline_resp.status_code}")
    print(f"  Projects found: {len(pipeline_resp.json())}")
    for p in pipeline_resp.json():
        print(f"  - {p['project_name']}: status={p['status']}, units={p['units']}")
    assert pipeline_resp.status_code == 200

    # Test delete document
    delete_resp = client.delete(f"/api/v1/data-bank/document/{doc_id}", headers=headers)
    print(f"\nDelete document {doc_id}: {delete_resp.status_code}")
    assert delete_resp.status_code == 204

    # Verify deletion
    verify_resp = client.get(f"/api/v1/data-bank/document/{doc_id}", headers=headers)
    print(f"Verify deleted: {verify_resp.status_code}")
    assert verify_resp.status_code == 404

    print("\n=== ALL API ENDPOINT TESTS PASSED ===")


if __name__ == "__main__":
    print("=" * 60)
    print("DATA BANK EXTRACTION SERVICE — TEST SUITE")
    print("=" * 60)

    print("\n--- Test 1: Document Type Detection ---")
    test_document_type_detection()

    print("\n--- Test 2: Extraction Service (Python fallback) ---")
    test_extraction_service()

    print("\n--- Test 3: API Endpoints ---")
    test_api_endpoints()

    print("\n" + "=" * 60)
    print("ALL TESTS PASSED")
    print("=" * 60)
